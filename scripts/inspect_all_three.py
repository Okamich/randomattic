# -*- coding: utf-8 -*-
import openpyxl, sys, json, os, re

sys.stdout.reconfigure(encoding='utf-8')

def extract_all():
    print("=== EXTRACTING NPC_JOB ===")
    wb_job = openpyxl.load_workbook('files/NPC/npc_job.xlsx', data_only=True)
    sheet_job = wb_job.active
    jobs = []
    for r in range(2, sheet_job.max_row + 1):
        num = sheet_job.cell(r, 1).value
        val = sheet_job.cell(r, 2).value
        if val is not None:
            jobs.append({"roll": num, "titleRu": str(val).strip()})
    print(f"Jobs extracted: {len(jobs)}")

    print("\n=== EXTRACTING NPC_VOICE ===")
    wb_voice = openpyxl.load_workbook('files/NPC/npc_voice.xlsx', data_only=True)
    sheet_voice = wb_voice.active
    # Col 1-2: Speed (Wait: let's verify what values are in col 2 and col 6)
    # Row 3: '1d3 Speed …' at C1, '1d3 ... and Pitch …' at C5, '1d8 Vocal Texture' at C9
    # Row 5-7:
    speed = []
    pitch = []
    texture = []
    for r in range(5, 8):
        c1, c2 = sheet_voice.cell(r, 1).value, sheet_voice.cell(r, 2).value
        c5, c6 = sheet_voice.cell(r, 5).value, sheet_voice.cell(r, 6).value
        if c1 is not None and c2 is not None:
            speed.append({"roll": c1, "val": str(c2).strip()})
        if c5 is not None and c6 is not None:
            pitch.append({"roll": c5, "val": str(c6).strip()})
    for r in range(5, 13):
        c9, c10 = sheet_voice.cell(r, 9).value, sheet_voice.cell(r, 10).value
        if c9 is not None and c10 is not None:
            texture.append({"roll": c9, "val": str(c10).strip()})
    
    # Col 1-2: 1d50 Speech patterns
    # Col 5-6: 1d50 Mannerisms
    speech_patterns = []
    mannerisms = []
    for r in range(16, sheet_voice.max_row + 1):
        c1, c2 = sheet_voice.cell(r, 1).value, sheet_voice.cell(r, 2).value
        c5, c6 = sheet_voice.cell(r, 5).value, sheet_voice.cell(r, 6).value
        if c1 is not None and c2 is not None:
            speech_patterns.append({"roll": c1, "textEn": str(c2).strip()})
        if c5 is not None and c6 is not None:
            mannerisms.append({"roll": c5, "textEn": str(c6).strip()})

    print(f"Speed: {speed}")
    print(f"Pitch: {pitch}")
    print(f"Texture: {len(texture)} items -> {[t['val'] for t in texture]}")
    print(f"Speech patterns: {len(speech_patterns)} items")
    print(f"Mannerisms: {len(mannerisms)} items")

    # Let's inspect NPC_TRAITS
    print("\n=== EXTRACTING NPC_TRAITS ===")
    wb_traits = openpyxl.load_workbook('files/NPC/NPC_TRAITS.xlsx', data_only=True)
    for sname in wb_traits.sheetnames:
        sheet = wb_traits[sname]
        print(f"\n--- Sheet: {sname} ---")
        for r in range(1, sheet.max_row + 1):
            row_items = []
            for c in range(1, sheet.max_column + 1):
                val = sheet.cell(r, c).value
                if val is not None:
                    row_items.append((c, str(val).strip()))
            if row_items:
                print(f"R{r}: {row_items}")

if __name__ == '__main__':
    extract_all()
