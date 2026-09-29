# -*- coding: utf-8 -*-
import openpyxl, sys, re, json
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'files/NPC/npc_self_motives.xlsx', data_only=True)
ws = wb['Мотивы и Слухи']

blocks = [
    (3, 10, ['Reaction', 'Motivation', 'Area']),
    (14, 22, ['A', 'B', 'C']),
    (24, 31, ['D', 'E', 'F']),
    (33, 40, ['G', 'H', 'I']),
    (42, 49, ['J', 'K', 'L']),
    (51, 58, ['M', 'N', 'O']),
    (60, 67, ['P', 'Q', 'R']),
    (69, 76, ['S', 'T', 'U']),
    (78, 85, ['V', 'W', 'X']),
    (87, 94, ['Y', 'Z', 'AA']),
    (96, 103, ['BB', 'CC'])
]

all_tables = {}
for r_start, r_end, tags in blocks:
    for c_idx, tag in enumerate(tags, start=1):
        header = ws.cell(r_start, c_idx).value or ''
        header = str(header).strip()
        rows = []
        
        # Special case for Z where first option was joined to header
        if tag == 'Z' and '1-3.' in header:
            parts = header.split('\xa0')
            if len(parts) == 1:
                parts = header.split(' 1-3.')
                header = parts[0].strip()
                rows.append('1-3.' + parts[1].strip())
            else:
                header = parts[0].strip()
                rows.append(parts[1].strip())
                
        for r in range(r_start + 1, r_end + 1):
            val = ws.cell(r, c_idx).value
            if val is not None:
                rows.append(str(val).strip())
        all_tables[tag] = {
            'header': header,
            'rows': rows
        }

print('Parsed total tables:', len(all_tables))
for k, v in all_tables.items():
    print(f'Table [{k}]: rows={len(v["rows"])} | header: {v["header"][:45]}...')

with open('files/NPC/motives_extracted_raw.json', 'w', encoding='utf-8') as f:
    json.dump(all_tables, f, ensure_ascii=False, indent=2)

print('Saved to files/NPC/motives_extracted_raw.json')
